import os
from pathlib import Path
from typing import Tuple

import numpy as np
import numpy.testing as npt
import pytest
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from standard_evaluator import (
    IntVariable,
    FloatVariable,
    ArrayVariable,
    Variable,
)
from standard_evaluator.evaluators import (
    get_interface_from_excel_named_ranges,
    create_variable_from_excel_address,
)

from standard_evaluator.evaluators.excel_utilities import (
    get_defined_names_info,
    get_cell_shape,
    _MAX_ROWS,
    _MAX_COLS,
    check_types,
    range_shape,
)


@pytest.fixture
def create_workbook(tmp_path):
    def _make(fn: str, setup_fn=None):
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        if setup_fn:
            setup_fn(ws)
        path = tmp_path / fn
        wb.save(path)
        return path

    return _make


def test_single_cell_input(create_workbook):
    def setup(ws):
        ws["A1"] = 123.0
        # create defined name referencing single cell
        dn = DefinedName("MY_INPUT", attr_text="Sheet1!$A$1")
        ws.parent.defined_names.add(dn)

    path = create_workbook("single_cell.xlsx", setup)
    interface, diagnostics = get_interface_from_excel_named_ranges(
        str(path), name="test1"
    )
    assert len(diagnostics) == 0
    assert any(v.name == "MY_INPUT" for v in interface.inputs)
    assert all(v.name != "MY_INPUT" for v in interface.outputs)


def test_range_array_and_formula(create_workbook):
    def setup(ws):
        # Fill a 1x3 row with numbers and put a formula in last cell
        ws["A1"] = 1
        ws["B1"] = 2
        ws["C1"] = "=SUM(A1:B1)"  # formula
        dn = DefinedName("MY_ROW", attr_text="Sheet1!$A$1:$C$1")
        ws.parent.defined_names.add(dn)

    path = create_workbook("row_range.xlsx", setup)
    interface, diagnostics = get_interface_from_excel_named_ranges(
        str(path), name="test2"
    )
    # It should be treated as an output because one cell is a formula
    assert any(v.name == "MY_ROW" for v in interface.outputs)


def test_nonexistent_sheet(create_workbook):
    # Create a workbook with a defined name referencing a missing sheet
    wb = Workbook()
    path = create_workbook("missing_sheet.xlsx")
    # Append a defined name manually on the saved workbook object
    # We'll reopen, add the defined name then save again
    from openpyxl import load_workbook

    wb2 = load_workbook(path)
    dn = DefinedName("BAD_REF", attr_text="NoSheet!$A$1")
    wb2.defined_names.add(dn)
    wb2.save(path)
    interface, diagnostics = get_interface_from_excel_named_ranges(
        str(path), name="test3"
    )
    # Diagnostics should include an error/warning about missing sheet
    assert any(d.get("name") == "BAD_REF" for d in diagnostics)


@pytest.mark.parametrize(
    "coord, expected",
    [
        ("A1", (True, (1, 1))),
        ("B2:B2", (True, (1, 1))),
        ("A1:B2", (False, (2, 2))),
        ("C3:C5", (False, (3, 1))),
        ("A1:D1", (False, (4,))),  # single row -> 1-D shape
        ("Z10:AA12", (False, (3, 2))),
        ("XFD1:XFD1", (True, (1, 1))),  # last Excel column single cell
    ],
)
def test_valid_shapes(coord: str, expected: Tuple[bool, Tuple[int, ...]]):
    result = get_cell_shape(coord)
    assert result == expected


def test_large_valid_range_at_limits():
    # max row and max column should be valid
    max_row = _MAX_ROWS
    max_col_letter = "XFD"  # 16384 in Excel column letters (XFD)
    coord = f"A1:{max_col_letter}{max_row}"
    single, shape = get_cell_shape(coord)
    assert single is False
    # rows = _MAX_ROWS, columns = _MAX_COLS
    assert shape == (_MAX_ROWS, _MAX_COLS)


@pytest.mark.parametrize(
    "bad_coord, exc_msg_fragment",
    [
        ("G12:A1", "minimum row/column greater than maximum"),
    ],
)
def test_invalid_min_greater_than_max(bad_coord: str, exc_msg_fragment: str):
    with pytest.raises(ValueError) as excinfo:
        get_cell_shape(bad_coord)
    assert exc_msg_fragment in str(excinfo.value)


def test_exceeds_row_limit():
    # construct a coordinate that exceeds the Excel row limit
    too_big_row = _MAX_ROWS + 1
    coord = f"A1:A{too_big_row}"
    with pytest.raises(ValueError) as excinfo:
        get_cell_shape(coord)
    assert "exceeds Excel limits" in str(excinfo.value)


def test_exceeds_col_limit():
    # construct a coordinate that exceeds the Excel column limit by using a numeric column index > _MAX_COLS
    # openpyxl accepts column letters only, so construct by converting number to column letters.
    # Here we produce a coordinate by using the last valid column + 1 in letter form.
    # For simplicity we reference the numeric boundary by using column index conversion via openpyxl if needed.
    from openpyxl.utils import get_column_letter

    too_big_col = _MAX_COLS + 1
    col_letter = get_column_letter(too_big_col)
    coord = f"{col_letter}1:{col_letter}1"
    with pytest.raises(ValueError) as excinfo:
        get_cell_shape(coord)
    assert "exceeds Excel limits" in str(excinfo.value)


def test_single_row_multiple_columns_returns_1d_shape():
    res = get_cell_shape("B1:E1")
    assert res == (False, (4,))


def test_single_column_multiple_rows_returns_2d_shape():
    res = get_cell_shape("C2:C6")
    assert res == (False, (5, 1))


@pytest.mark.parametrize(
    "input_val, expected",
    [
        (int, "int"),
        (float, "float"),
        ("int", "int"),
        ("float", "float"),
    ],
)
def test_valid_inputs_return_normalized_string(input_val, expected):
    assert check_types(input_val) == expected


def test_default_none_rejected():
    with pytest.raises(TypeError):
        check_types(None)


@pytest.mark.parametrize("bad", [str, list, dict, "integer", "FLOAT", 123])
def test_invalid_types_raise_TypeError(bad):
    with pytest.raises(TypeError):
        check_types(bad)


def test_create_single_cell_int_by_type_object():
    v = create_variable_from_excel_address("v1", "Sheet1!A1", var_type=int)
    assert isinstance(v, IntVariable)
    assert getattr(v, "name", None) == "v1"
    assert v.options.get("ExcelLocation") == "Sheet1!A1"
    # class_type is passed through in the implementation; expect normalized string "int"
    assert getattr(v, "class_type", None) == "int"


def test_create_single_cell_float_by_string():
    v = create_variable_from_excel_address("v2", "SheetA!B2", var_type="float")
    assert isinstance(v, FloatVariable)
    assert v.name == "v2"
    assert v.options["ExcelLocation"] == "SheetA!B2"
    assert v.class_type == "float"


def test_create_array_variable_from_range():
    v = create_variable_from_excel_address("arr", "S!C3:D4", var_type=float)
    assert isinstance(v, ArrayVariable)
    assert v.name == "arr"
    assert v.options["ExcelLocation"] == "S!C3:D4"
    # C3:D4 is 2 rows x 2 cols -> shape should be (2, 2)
    assert getattr(v, "shape", None) == (2, 2)


def test_single_row_range_returns_1d_shape():
    v = create_variable_from_excel_address("row", "Sheet1!A1:D1", var_type=float)
    assert isinstance(v, ArrayVariable)
    # A1:D1 is single row 4 columns -> shape should be (4,)
    assert v.shape == (4,)


@pytest.mark.parametrize(
    "bad_address",
    [
        "Sheet1A1",  # missing '!'
        "S!A1!B2",  # multiple '!'
        "!A1",  # empty sheet name
        "Sheet1!",  # empty coord
        "Sheet1!A0",  # invalid row 0
        f"Sheet1!A{_MAX_ROWS+1}",  # row exceeds Excel limit
    ],
)
def test_invalid_addresses_raise_value_error(bad_address):
    with pytest.raises(ValueError):
        create_variable_from_excel_address("x", bad_address, var_type=float)


def test_invalid_var_type_raises_type_error_for_single_cell():
    # Using a single-cell address so check_types is exercised
    with pytest.raises(TypeError):
        create_variable_from_excel_address("x", "Sheet1!A1", var_type="integer")

# --- Default value tests ---


def test_default_value_scalar_input(create_workbook):
    """Scalar input cell should have its value stored as default."""
    def setup(ws):
        ws["A1"] = 42.5
        dn = DefinedName("MY_SCALAR", attr_text="Sheet1!$A$1")
        ws.parent.defined_names.add(dn)

    path = create_workbook("scalar_default.xlsx", setup)
    interface, _ = get_interface_from_excel_named_ranges(str(path), name="test_def")
    var = next(v for v in interface.inputs if v.name == "MY_SCALAR")
    assert var.default == 42.5


def test_default_value_array_input(create_workbook):
    """Array input range should have its values stored as a numpy default."""
    def setup(ws):
        ws["A1"] = 1.0
        ws["B1"] = 2.0
        ws["C1"] = 3.0
        dn = DefinedName("MY_ROW", attr_text="Sheet1!$A$1:$C$1")
        ws.parent.defined_names.add(dn)

    path = create_workbook("array_default.xlsx", setup)
    interface, _ = get_interface_from_excel_named_ranges(str(path), name="test_def")
    var = next(v for v in interface.inputs if v.name == "MY_ROW")
    assert var.default is not None
    npt.assert_array_equal(var.default, np.array([1.0, 2.0, 3.0]))


def test_default_value_2d_array(create_workbook):
    """2D array range should have its values stored as a 2D numpy default."""
    def setup(ws):
        ws["A1"] = 1.0
        ws["B1"] = 2.0
        ws["A2"] = 3.0
        ws["B2"] = 4.0
        dn = DefinedName("MY_MATRIX", attr_text="Sheet1!$A$1:$B$2")
        ws.parent.defined_names.add(dn)

    path = create_workbook("matrix_default.xlsx", setup)
    interface, _ = get_interface_from_excel_named_ranges(str(path), name="test_def")
    var = next(v for v in interface.inputs if v.name == "MY_MATRIX")
    assert var.default is not None
    npt.assert_array_equal(var.default, np.array([[1.0, 2.0], [3.0, 4.0]]))


def test_default_value_formula_output(create_workbook):
    """Formula (output) cell should also have its cached value as default."""
    def setup(ws):
        ws["A1"] = 10.0
        ws["B1"] = "=A1*2"
        dn = DefinedName("MY_FORMULA", attr_text="Sheet1!$B$1")
        ws.parent.defined_names.add(dn)

    path = create_workbook("formula_default.xlsx", setup)
    interface, _ = get_interface_from_excel_named_ranges(str(path), name="test_def")
    var = next(v for v in interface.outputs if v.name == "MY_FORMULA")
    # openpyxl data_only on a freshly-saved workbook may return None for
    # formulas that have never been evaluated by Excel, so accept None here.
    assert var.default is None or isinstance(var.default, float)


def test_default_value_empty_cell_is_none(create_workbook):
    """An empty cell should result in default=None."""
    def setup(ws):
        # A1 is left empty
        dn = DefinedName("MY_EMPTY", attr_text="Sheet1!$A$1")
        ws.parent.defined_names.add(dn)

    path = create_workbook("empty_default.xlsx", setup)
    interface, _ = get_interface_from_excel_named_ranges(str(path), name="test_def")
    var = next(v for v in interface.inputs if v.name == "MY_EMPTY")
    assert var.default is None


def test_range_shape_basic():
    assert range_shape("A1:A1") == (1, 1)
    assert range_shape("A3:B28") == (26, 2)
    assert range_shape("C5:E10") == (6, 3)


def test_range_shape_with_dollar_signs():
    assert range_shape("$A$3:$B$28") == (26, 2)
    assert range_shape("C$5:$E10") == (6, 3)


def test_range_shape_invalid_format():
    with pytest.raises(ValueError):
        range_shape("InvalidRange")
    with pytest.raises(ValueError):
        range_shape("A1B2")
    with pytest.raises(ValueError):
        range_shape("A1:A")  # Missing row number


def test_range_shape_zero_or_negative():
    with pytest.raises(ValueError):
        range_shape("B5:A3")  # end row < start row
    with pytest.raises(ValueError):
        range_shape("C10:C5")  # end row < start row
    with pytest.raises(ValueError):
        range_shape("B2:A2")  # end col < start col
